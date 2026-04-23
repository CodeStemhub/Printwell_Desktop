"""
OCR Service - Extracts text from documents.

WHAT IT IS: The text extraction engine.
WHY IT EXISTS: The AI cannot read a file — it reads text.
              This converts every document type into text the LLM processes.
WHAT IT DOES:
    Routes each document to the correct extractor based on file type.
    Low OCR confidence triggers Google Vision as a smarter fallback.
"""

import io
from typing import Dict, Any, Optional
from PIL import Image
import pytesseract
from sqlalchemy.orm import Session

from config import settings


def extract_from_pdf(file_path: str) -> Dict[str, Any]:
    """
    Extract text and tables from PDF files.
    
    WHY pdfplumber: Preserves table structure — critical for financial
    documents where numbers must stay in correct columns.
    """
    import pdfplumber
    
    with pdfplumber.open(file_path) as pdf:
        full_text = ""
        tables = []
        
        for page in pdf.pages:
            text = page.extract_text()
            if text:
                full_text += text + "\n"
            
            page_tables = page.extract_tables()
            if page_tables:
                tables.extend(page_tables)
    
    return {
        "text": full_text,
        "tables": tables,
        "method": "pdfplumber",
        "confidence": 0.95
    }


def extract_from_image(file_path: str) -> Dict[str, Any]:
    """
    Extract text from images using Tesseract OCR.
    
    WHY two-step: Tesseract is free and fast but struggles with
    low-quality scans. Google Vision costs money but handles anything.
    We try free first, pay only when needed.
    """
    img = Image.open(file_path)
    
    # Configure Tesseract for better accuracy
    tesseract_config = '--oem 3 --psm 6'
    text = pytesseract.image_to_string(img, config=tesseract_config)
    
    # Get confidence scores
    data = pytesseract.image_to_data(
        img, 
        output_type=pytesseract.Output.DICT
    )
    
    confidences = [c for c in data['conf'] if c > 0]
    confidence_score = (sum(confidences) / len(confidences) / 100) if confidences else 0
    
    # If confidence is low, use Google Vision as fallback
    if confidence_score < 0.75 and settings.GOOGLE_VISION_API_KEY:
        return extract_with_google_vision(file_path)
    
    return {
        "text": text,
        "tables": [],
        "method": "tesseract",
        "confidence": confidence_score
    }


def extract_with_google_vision(file_path: str) -> Dict[str, Any]:
    """
    Fallback OCR using Google Vision API for low-quality images.
    
    More expensive but handles handwriting, poor lighting, and complex layouts better.
    """
    from google.cloud import vision
    
    client = vision.ImageAnnotatorClient(
        credentials=settings.GOOGLE_VISION_API_KEY
    )
    
    with io.open(file_path, 'rb') as image_file:
        content = image_file.read()
    
    image = vision.Image(content=content)
    response = client.text_detection(image=image)
    texts = response.text_annotations
    
    if texts:
        text = texts[0].description
        confidence = 0.85  # Google Vision typically has high confidence
    else:
        text = ""
        confidence = 0.0
    
    return {
        "text": text,
        "tables": [],
        "method": "google_vision",
        "confidence": confidence
    }


def extract_from_docx(file_path: str) -> Dict[str, Any]:
    """
    Extract text from Word documents.
    """
    from docx import Document
    
    doc = Document(file_path)
    full_text = []
    
    for para in doc.paragraphs:
        if para.text.strip():
            full_text.append(para.text)
    
    # Also extract text from tables
    tables = []
    for table in doc.tables:
        table_data = []
        for row in table.rows:
            row_data = [cell.text for cell in row.cells]
            table_data.append(row_data)
        if table_data:
            tables.append(table_data)
    
    return {
        "text": "\n".join(full_text),
        "tables": tables,
        "method": "python-docx",
        "confidence": 0.98
    }


def extract_from_xlsx(file_path: str) -> Dict[str, Any]:
    """
    Extract data from Excel spreadsheets.
    """
    from openpyxl import load_workbook
    
    wb = load_workbook(file_path, data_only=True)
    full_text = []
    all_tables = []
    
    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        full_text.append(f"=== Sheet: {sheet_name} ===\n")
        
        sheet_data = []
        for row in sheet.iter_rows(values_only=True):
            if any(cell is not None for cell in row):
                row_str = " | ".join(str(cell) if cell is not None else "" for cell in row)
                full_text.append(row_str)
                sheet_data.append(list(row))
        
        if sheet_data:
            all_tables.append({"sheet": sheet_name, "data": sheet_data})
    
    return {
        "text": "\n".join(full_text),
        "tables": all_tables,
        "method": "openpyxl",
        "confidence": 0.98
    }


async def extract_text(document) -> Dict[str, Any]:
    """
    Main entry point - routes to correct extractor based on content type.
    
    Args:
        document: Document model instance
        
    Returns:
        Dictionary with extracted text, tables, method used, and confidence score
    """
    content_type = document.content_type
    
    # For local development, we'll use storage_url as local path
    # In production, download from Supabase first
    file_path = document.storage_url
    
    if content_type == "application/pdf":
        return extract_from_pdf(file_path)
    
    elif content_type in ["image/jpeg", "image/png", "image/webp"]:
        return extract_from_image(file_path)
    
    elif "wordprocessingml" in content_type:
        return extract_from_docx(file_path)
    
    elif "spreadsheetml" in content_type:
        return extract_from_xlsx(file_path)
    
    else:
        raise ValueError(f"Unsupported content type: {content_type}")
